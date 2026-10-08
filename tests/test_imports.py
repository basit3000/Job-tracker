"""Fictional imports through real parsers, routes, transactions and sync."""

import io
import json
import re
from datetime import date, timedelta
from zipfile import ZipFile

import openpyxl
import pytest
import requests
import xlwt

from app.contracts import ServiceError
from app.extensions import db
from app.google_sheets import (
    read_private_sheet,
    read_public_sheet,
    sheet_reference,
)
from app.import_readers import MAX_BYTES, read_file
from app.import_service import (
    commit_batch,
    create_batch,
    initial_options,
    options_digest,
    preview,
    save_options,
)
from app.models import ChangeEntry, ImportBatch, JobApplication, _utcnow
from tests.google_provider import (
    BASE_URL,
    GOOGLE_CONFIG,
    mock_google_transport,
)
from tests.helpers import csrf_token

CSV = (
    b"Position;Employer;Application stage;Date applied;Job li"
    b"nk\nEngineer;Example;Submitted;06/10/2026;https://exampl"
    b"e.com/jobs/1\n"
)


def upload(client, data=CSV, filename="jobs.csv"):
    return client.post(
        "/imports",
        data={
            "source": "file",
            "file": (io.BytesIO(data), filename),
            "csrf_token": csrf_token(client, "/imports"),
        },
    )


def options_form(options, token):
    values = {
        key: str(options[key])
        for key in (
            "sheet",
            "header",
            "date_order",
            "default_status",
            "duplicates",
        )
    }
    values.update(
        {
            f"column_{index}": field
            for index, field in enumerate(options["mapping"])
        }
    )
    values.update(
        {
            "status_" + key: value
            for key, value in options["status_map"].items()
        }
    )
    for key in ("preserve", "skip_invalid"):
        if options[key]:
            values[key] = "yes"
    values["csrf_token"] = token
    return values


def test_import_preview_confirm_replay_sync_and_source_erasure(
    app, logged_client, users
):
    result = upload(logged_client)
    assert result.status_code == 302
    map_url = result.location
    assert logged_client.get(map_url).status_code == 200
    with app.app_context():
        batch = ImportBatch.query.one()
        batch_id = batch.id
        options = initial_options(batch)
        assert options["mapping"] == [
            "title",
            "company",
            "status",
            "appliedDate",
            "url",
        ]
        options["date_order"] = "dmy"
    result = logged_client.post(
        map_url, data=options_form(options, csrf_token(logged_client, map_url))
    )
    assert result.status_code == 302
    review_url = result.location
    review = logged_client.get(review_url)
    assert b"Engineer" in review.data and b"2026-10-06" in review.data
    digest = re.search(rb'name="digest" value="([^"]+)"', review.data)[
        1
    ].decode()
    for _ in range(2):
        assert (
            logged_client.post(
                f"/imports/{batch_id}/confirm",
                data={
                    "digest": digest,
                    "csrf_token": csrf_token(logged_client, review_url),
                },
            ).status_code
            == 302
        )
    with app.app_context():
        job = JobApplication.query.one()
        assert job.user_id == users[0] and job.status == "applied"
        assert job.applied_on == date(2026, 10, 6)
        assert job.history[0].source == "import"
        assert ChangeEntry.query.one().source == "import"
        batch = db.session.get(ImportBatch, batch_id)
        assert batch.payload is None and batch.options is None
        assert batch.result["imported"] == 1


@pytest.mark.parametrize(
    "filename,data",
    [
        ("jobs.csv", b"Role,Organization\nEngineer,Example\n"),
        ("jobs.tsv", b"Role\tOrganization\nEngineer\tExample\n"),
        ("jobs.txt", b"Role|Organization\nEngineer|Example\n"),
        (
            "jobs.json",
            b'{"applications":[{"title":"Engineer","company":"Example"}]}',
        ),
        ("jobs.jsonl", b'{"role":"Engineer","employer":"Example"}\n'),
        ("jobs.ndjson", b'{"role":"Engineer","employer":"Example"}\n'),
        ("jobs.csv", "Role;Organization\nEngineer;Example\n".encode("utf-16")),
    ],
)
def test_text_formats_and_encodings(filename, data):
    sheets = read_file(data, filename)
    assert sheets[0]["rows"][1] == ["Engineer", "Example"]


def test_xlsx_tabs_preamble_dates_and_formula_handling():
    book = openpyxl.Workbook()
    first = book.active
    first.title = "Notes"
    first.append(["Instructions only"])
    sheet = book.create_sheet("Applications")
    sheet.append(["My tracker"])
    sheet.append(["Position", "Employer", "Applied on", "Comments"])
    sheet.append(
        [
            "Engineer",
            "Example",
            date(2026, 10, 6),
            '=HYPERLINK("https://evil.example.com")',
        ]
    )
    stream = io.BytesIO()
    book.save(stream)
    tables = read_file(stream.getvalue(), "jobs.xlsx")
    assert [item["name"] for item in tables] == ["Notes", "Applications"]
    assert tables[1]["rows"][2] == ["Engineer", "Example", "2026-10-06", ""]


def test_legacy_xls_has_real_excel_dates():
    book = xlwt.Workbook()
    sheet = book.add_sheet("Applications")
    for column, value in enumerate(["Role", "Employer", "Applied on"]):
        sheet.write(0, column, value)
    sheet.write(1, 0, "Engineer")
    sheet.write(1, 1, "Example")
    sheet.write(
        1, 2, date(2026, 10, 6), xlwt.easyxf(num_format_str="DD/MM/YYYY")
    )
    stream = io.BytesIO()
    book.save(stream)
    assert read_file(stream.getvalue(), "jobs.xls")[0]["rows"][1] == [
        "Engineer",
        "Example",
        "2026-10-06",
    ]


def test_ods_reads_repeated_cells_and_dates():
    document = (
        b'<office:document-content xmlns:office="urn:oasis:names:'
        b'tc:opendocument:xmlns:office:1.0" xmlns:table="urn:oasi'
        b's:names:tc:opendocument:xmlns:table:1.0" xmlns:text="ur'
        b'n:oasis:names:tc:opendocument:xmlns:text:1.0"><office:b'
        b'ody><office:spreadsheet><table:table table:name="Applic'
        b'ations"><table:table-row><table:table-cell><text:p>Role'
        b"</text:p></table:table-cell><table:table-cell><text:p>E"
        b"mployer</text:p></table:table-cell></table:table-row><t"
        b"able:table-row><table:table-cell><text:p>Engineer</text"
        b":p></table:table-cell><table:table-cell><text:p>Example"
        b"</text:p></table:table-cell><table:table-cell table:num"
        b'ber-columns-repeated="1024"/></table:table-row></table:'
        b"table></office:spreadsheet></office:body></office:docum"
        b"ent-content>"
    )
    stream = io.BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("content.xml", document)
    assert read_file(stream.getvalue(), "jobs.ods")[0]["rows"][1] == [
        "Engineer",
        "Example",
    ]


@pytest.mark.parametrize(
    "data,filename",
    [
        (b"broken", "a.xlsx"),
        (b"{", "a.json"),
        (b"broken", "a.exe"),
        (b"", "a.csv"),
        (b"a" * (MAX_BYTES + 1), "a.csv"),
        (b'{"applications":[]}', "a.json"),
    ],
    ids=[
        "bad-xlsx",
        "bad-json",
        "unsupported",
        "empty",
        "oversized",
        "empty-records",
    ],
)
def test_bad_or_oversized_files_are_rejected(data, filename):
    with pytest.raises(ServiceError):
        read_file(data, filename)


def test_xml_entities_and_archive_expansion_are_rejected():
    document = (
        b'<!DOCTYPE x [<!ENTITY private SYSTEM "file:///etc/passw'
        b'd">]><x>&private;</x>'
    )
    stream = io.BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("content.xml", document)
    with pytest.raises(ServiceError):
        read_file(stream.getvalue(), "a.ods")


def test_row_column_and_cell_limits_keep_actionable_messages():
    for data, message in (
        (b"a,b\n" * 1031, "1,000"),
        ((",".join(["cell"] * 61)).encode(), "60 columns"),
        (b"a" * 20001, "20,000"),
    ):
        with pytest.raises(ServiceError, match=message):
            read_file(data, "jobs.csv")


def test_mapping_ambiguous_dates_unknown_statuses_and_duplicates(app, users):
    data = (
        b"Tracker title\nJob_title,Company name,Applied on,Stage,E"
        b"xtra\nEngineer,Example,06/10/2026,Waiting for recruiter,"
        b"Keep this\nEngineer,Example,06/10/2026,Waiting for recru"
        b"iter,Keep this\n"
    )
    with app.app_context():
        batch = create_batch(users[0], read_file(data, "jobs.csv"))
        options = initial_options(batch)
        assert options["header"] == 2
        rows, counts = preview(batch, options)
        assert counts["invalid"] == 2 and "Ambiguous" in rows[0]["error"]
        options.update(date_order="dmy", preserve=True)
        rows, counts = preview(batch, options)
        assert "Unknown status" in rows[0]["error"]
        options["status_map"] = {"waitingforrecruiter": "applied"}
        save_options(batch, options)
        rows, counts = preview(batch, options)
        assert counts == {"ready": 1, "invalid": 0, "duplicates": 1}
        assert "Extra: Keep this" in rows[0]["fields"]["note"]
        commit_batch(users[0], batch.id, options_digest(options))
        repeated = create_batch(users[0], read_file(data, "jobs.csv"))
        assert preview(repeated, options)[1]["ready"] == 0
        other = create_batch(users[1], read_file(data, "jobs.csv"))
        assert preview(other, options)[1]["ready"] == 1


def test_invalid_rows_require_explicit_skip_and_commit_is_atomic(
    app, users, monkeypatch
):
    from app import import_service

    with app.app_context():
        batch = create_batch(
            users[0],
            read_file(
                (
                    b"Role,Employer,Link\nFirst,Example,https://example.com/1\n"
                    b"Bad,Example,javascript:evil\nSecond,Example,https://exam"
                    b"ple.com/2\n"
                ),
                "jobs.csv",
            ),
        )
        options = initial_options(batch)
        save_options(batch, options)
        with pytest.raises(ServiceError, match="Resolve invalid"):
            commit_batch(users[0], batch.id, options_digest(options))
        assert JobApplication.query.count() == 0
        options["skip_invalid"] = True
        save_options(batch, options)
        original = import_service.write_application
        calls = []

        def fail_second(*args, **kwargs):
            calls.append(1)
            if len(calls) == 2:
                raise RuntimeError("fictional failure")
            return original(*args, **kwargs)

        monkeypatch.setattr(import_service, "write_application", fail_second)
        with pytest.raises(RuntimeError):
            commit_batch(users[0], batch.id, options_digest(options))
        assert JobApplication.query.count() == ChangeEntry.query.count() == 0
        assert db.session.get(ImportBatch, batch.id).payload is not None
        monkeypatch.setattr(import_service, "write_application", original)
        assert (
            commit_batch(users[0], batch.id, options_digest(options))[
                "imported"
            ]
            == 2
        )


def test_ownership_csrf_expiry_cancel_and_stale_preview(
    app, logged_client, users
):
    client = app.test_client()
    assert client.get("/imports").status_code == 302
    assert (
        logged_client.post(
            "/imports", data={"source": "paste", "text": "Title,Company\na,b"}
        ).status_code
        == 400
    )
    response = upload(logged_client)
    batch_id = response.location.split("/")[2]
    with app.app_context():
        batch = db.session.get(ImportBatch, batch_id)
        options = initial_options(batch)
        options["date_order"] = "dmy"
        save_options(batch, options)
        with pytest.raises(ServiceError, match="latest preview"):
            commit_batch(users[0], batch_id, "stale-digest")
    with client.session_transaction() as session:
        session["_user_id"] = str(users[1])
    for suffix in ("map", "review", "report.csv"):
        assert client.get(f"/imports/{batch_id}/{suffix}").status_code == 404
    with app.app_context():
        batch = db.session.get(ImportBatch, batch_id)
        batch.expires_at = _utcnow() - timedelta(seconds=1)
        db.session.commit()
    assert logged_client.get(response.location).status_code == 410
    assert (
        app.test_cli_runner().invoke(args=["cleanup-imports"]).exit_code == 0
    )
    with app.app_context():
        assert ImportBatch.query.count() == 0
    response = upload(logged_client)
    batch_id = response.location.split("/")[2]
    assert (
        logged_client.post(
            f"/imports/{batch_id}/cancel",
            data={"csrf_token": csrf_token(logged_client, response.location)},
        ).status_code
        == 302
    )
    with app.app_context():
        assert ImportBatch.query.count() == 0


def google_response(data, status=200, **headers):
    response = requests.Response()
    response.status_code = status
    response.headers.update(headers)
    response.raw = io.BytesIO(data)
    return response


@pytest.mark.parametrize(
    "url",
    [
        "http://docs.google.com/spreadsheets/d/fictionalSheet12345/edit",
        "https://localhost/private",
        (
            "https://docs.google.com.evil.example.com/spreadsheets/d"
            "/fictionalSheet12345/edit"
        ),
        (
            "https://user:password@docs.google.com/spreadsheets/d/fi"
            "ctionalSheet12345/edit"
        ),
        (
            "https://docs.google.com:invalid/spreadsheets/d/fictiona"
            "lSheet12345/edit"
        ),
    ],
)
def test_sheet_urls_cannot_fetch_arbitrary_servers(url):
    with pytest.raises(ServiceError):
        sheet_reference(url)


def test_google_public_redirect_controls_and_bounded_read(monkeypatch):
    reference = sheet_reference(
        (
            "https://docs.google.com/spreadsheets/d/fictionalSheet12"
            "345/edit#gid=42"
        )
    )
    calls = []

    def fake_get(url, **kwargs):
        calls.append((url, kwargs))
        return google_response(
            b"Role,Employer\nEngineer,Example", **{"Content-Type": "text/csv"}
        )

    monkeypatch.setattr(requests, "get", fake_get)
    assert read_public_sheet(reference)[0]["rows"][1] == [
        "Engineer",
        "Example",
    ]
    assert calls[0][1]["params"]["gid"] == "42"
    assert calls[0][1]["allow_redirects"] is False
    monkeypatch.setattr(
        requests,
        "get",
        lambda *a, **k: google_response(
            b"", 302, Location="http://127.0.0.1/private"
        ),
    )
    with pytest.raises(ServiceError, match="redirect"):
        read_public_sheet(reference)
    monkeypatch.setattr(
        requests,
        "get",
        lambda *a, **k: google_response(b"a" * (MAX_BYTES + 1)),
    )
    with pytest.raises(ServiceError, match="exceeds"):
        read_public_sheet(reference)


@pytest.mark.parametrize("source", ["private", "google"])
def test_private_google_authorization_is_optional_account_bound_and_transient(
    app_factory, monkeypatch, source
):
    app = app_factory(**GOOGLE_CONFIG)
    with app.app_context():
        from app.models import User

        user = User(email="fictional-importer@example.com")
        user.set_password("fictional-password")
        db.session.add(user)
        db.session.commit()
        user_id = user.id
    client = app.test_client()
    with client.session_transaction() as session:
        session["_user_id"] = str(user_id)
    provider = mock_google_transport(monkeypatch)
    response = client.post(
        "/imports",
        base_url=BASE_URL,
        data={
            "source": source,
            "sheet_access": "private",
            "import_mode": "once",
            "intent": "signin",
            "sheet_url": (
                "https://docs.google.com/spreadsheets/d/fictionalSheet12"
                "345/edit#gid=42"
            ),
            "csrf_token": csrf_token(client, "/imports"),
        },
    )
    from urllib.parse import parse_qs, urlsplit

    params = {
        key: value[0]
        for key, value in parse_qs(urlsplit(response.location).query).items()
    }
    assert "spreadsheets.readonly" in params["scope"]
    code = provider.issue(
        params["nonce"], code_challenge=params["code_challenge"]
    )

    def sheets_get(url, **kwargs):
        assert (
            kwargs["headers"]["Authorization"]
            == "Bearer fictional-access-token"
        )
        assert kwargs["allow_redirects"] is False
        payload = (
            {"values": [["Role", "Employer"], ["Engineer", "Example"]]}
            if "/values/" in url
            else {
                "sheets": [
                    {"properties": {"sheetId": 42, "title": "Applications"}}
                ]
            }
        )
        return google_response(json.dumps(payload).encode())

    monkeypatch.setattr(requests, "get", sheets_get)
    response = client.get(
        "/auth/google/callback",
        base_url=BASE_URL,
        query_string={"code": code, "state": params["state"]},
    )
    assert response.location.startswith("/imports/")
    with client.session_transaction() as session:
        assert session["_user_id"] == str(user_id)
        assert "fictional-access-token" not in json.dumps(dict(session))
    with app.app_context():
        batch = ImportBatch.query.one()
        assert batch.user_id == user_id
        assert "fictional-access-token" not in json.dumps(batch.payload)


def test_pasted_table_preview_report_formula_escaping_and_xss(
    app, logged_client
):
    response = logged_client.post(
        "/imports",
        data={
            "source": "paste",
            "text": (
                "Role\tEmployer\tNotes\n"
                "=1+1\tExample\t<script>alert(1)</script>"
            ),
            "csrf_token": csrf_token(logged_client, "/imports"),
        },
    )
    with app.app_context():
        batch = ImportBatch.query.one()
        options = initial_options(batch)
        save_options(batch, options)
        batch_id = batch.id
    report = logged_client.get(f"/imports/{batch_id}/report.csv")
    assert b"'=1+1" in report.data
    page = logged_client.get(response.location)
    assert b"<script>alert(1)</script>" not in page.data
    assert b"&lt;script&gt;" in page.data


def test_import_feed_matches_published_api_schema(app, users, client, paired):
    from tests.test_api_schema import validator

    with app.app_context():
        batch = create_batch(
            users[0], read_file(b"Role,Employer\nEngineer,Example", "jobs.csv")
        )
        options = initial_options(batch)
        save_options(batch, options)
        commit_batch(users[0], batch.id, options_digest(options))
    document = client.get("/api/v1/openapi.json").json
    for path in ("/api/v1/applications", "/api/v1/changes"):
        response = client.get(path, headers=paired)
        assert response.status_code == 200
        validator(document, "FeedPage").validate(response.json)
        assert response.json["entries"][0]["source"] == "import"


def test_private_sheet_request_never_forwards_token_on_redirect(monkeypatch):
    reference = sheet_reference(
        "https://docs.google.com/spreadsheets/d/fictionalSheet12345/edit"
    )
    calls = []

    def fake_get(url, **kwargs):
        calls.append((url, kwargs))
        return google_response(b"", 302, Location="https://evil.example.com/")

    monkeypatch.setattr(requests, "get", fake_get)
    with pytest.raises(ServiceError):
        read_private_sheet(reference, "fictional-access-token")
    assert len(calls) == 1
    assert calls[0][0].startswith("https://sheets.googleapis.com/")
    assert calls[0][1]["allow_redirects"] is False


def test_private_sheet_flow_cannot_switch_tracker_account(
    app_factory, monkeypatch
):
    from urllib.parse import parse_qs, urlsplit

    from app.models import User

    app = app_factory(**GOOGLE_CONFIG)
    with app.app_context():
        people = [
            User(email=f"person-{index}@example.com") for index in range(2)
        ]
        for person in people:
            person.set_password("fictional-password")
        db.session.add_all(people)
        db.session.commit()
        user_ids = [person.id for person in people]
    client = app.test_client()
    with client.session_transaction() as session:
        session["_user_id"] = str(user_ids[0])
    provider = mock_google_transport(monkeypatch)
    response = client.post(
        "/imports",
        base_url=BASE_URL,
        data={
            "source": "private",
            "sheet_url": "https://docs.google.com/spreadsheets/d/fictionalSheet12345/edit",
            "csrf_token": csrf_token(client, "/imports"),
        },
    )
    params = {
        key: value[0]
        for key, value in parse_qs(urlsplit(response.location).query).items()
    }
    code = provider.issue(
        params["nonce"], code_challenge=params["code_challenge"]
    )
    with client.session_transaction() as session:
        session["_user_id"] = str(user_ids[1])
    response = client.get(
        "/auth/google/callback",
        base_url=BASE_URL,
        query_string={"code": code, "state": params["state"]},
    )
    assert response.location == "/login"
    assert provider.token_calls == 0
    with app.app_context():
        assert ImportBatch.query.count() == 0
